"""Modèle Excel pour reconstituer l'historique (une ligne = une date × une ZS).
Usage : python3 scripts/creer-modele-historique.py [sortie.xlsx] [ZS1,ZS2,...]
Les en-têtes viennent de web/src/domain/historique.js (HIST_HEADERS) : lancés via node pour rester identiques à l'app."""
import json, subprocess, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(root, 'web/src/assets/SurvEpi_modele_historique.xlsx')
ZS = sys.argv[2].split(',') if len(sys.argv) > 2 else ['Isiro', 'Wamba', 'Pawa', 'Boma-Mangbetu', 'Dungu', 'Rungu', 'Gombari', 'Faradje']
js = "import('./src/domain/historique.js').then(m=>{const r=m.HIST_HEADERS;import('./src/domain/report.js').then(p=>console.log(JSON.stringify({h:r,req:p.FIELDS.filter(f=>f.required).map(f=>f.key),help:Object.fromEntries(p.FIELDS.filter(f=>f.help).map(f=>[f.key,f.help]))})))})"
data = json.loads(subprocess.check_output(['node', '--input-type=module', '-e', js], cwd=os.path.join(root, 'web')))
H, REQ, HELP = data['h'], set(data['req']), data['help']
N = 3000
DEEP, BLUE = '004E80', '0086C4'
thin = Side(style='thin', color='CFDFEC'); grid = Border(left=thin, right=thin, top=thin, bottom=thin)
GROUPES = {'alertes_validees': 'D9EEF8', 'ct_report': 'E8F4EA', 'ct_sortie_21j': 'FDF1DC', 'ct_vus': 'EFE8F6', 'reco_actifs': 'F3F3F3'}

wb = Workbook()
lire = wb.active; lire.title = 'À lire'
texte = [
    ('SurvEpi : reconstituer l’historique', 'title'), ('', None),
    ('Une ligne = une date et une zone de santé, avec les chiffres de la fiche journalière de ce jour-là.', None),
    ('Sources possibles : compilations journalières, rapports de situation (SitRep), messages WhatsApp des ZS.', None),
    ('Indiquer la source de chaque ligne dans la dernière colonne (ex. « SitRep n°12 », « Rapport ZS Wamba »).', None), ('', None),
    ('Ce qu’il faut savoir', 'h'),
    ('Une case vide = chiffre inconnu (pas zéro). Mettre 0 seulement quand c’est vraiment 0.', None),
    ('Les reports (investigations en attente, contacts à suivre, non-vus de la veille) se calculent tout seuls à partir du jour précédent : on peut laisser ces colonnes vides.', None),
    ('Les lignes incohérentes (ex. plus d’investigations faites que d’alertes) sont signalées à l’import et ne sont pas enregistrées : les corriger puis recharger le fichier.', None),
    ('Une fiche déjà saisie dans l’app n’est jamais écrasée. Recharger le fichier corrigé remplace l’historique déjà importé pour les mêmes dates.', None),
    ('Si les outils Excel DPS (investigations, suivi des contacts) sont remplis, les charger aussi à l’import : l’app en tire les chiffres manquants. Ce fichier-ci reste prioritaire.', None),
    ('', None), ('Colonnes', 'h'),
]
for k, label in H[2:-2]:
    texte.append((f'{label}{" (utile pour la complétude)" if k in REQ else ""}' + (f' : {HELP[k]}' if k in HELP else ''), None))
for i, (t, kind) in enumerate(texte, 1):
    c = lire.cell(row=i, column=1, value=t)
    if kind == 'title': c.font = Font(bold=True, size=16, color=DEEP)
    elif kind == 'h': c.font = Font(bold=True, size=12, color=BLUE)
    c.alignment = Alignment(wrap_text=True, vertical='top')
lire.column_dimensions['A'].width = 140; lire.sheet_view.showGridLines = False

ws = wb.create_sheet('Historique')
fill = None
for j, (k, label) in enumerate(H, 1):
    name = label + (' *' if k in ('date', 'zs') else '')
    c = ws.cell(row=1, column=j, value=name)
    c.font = Font(bold=True, color='FFFFFF')
    if k in GROUPES: fill = GROUPES[k]
    c.fill = PatternFill('solid', fgColor='B3302A' if k in ('date', 'zs') else DEEP)
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = grid
    if k in HELP: c.comment = Comment(HELP[k], 'SurvEpi')
    ws.column_dimensions[c.column_letter].width = 13 if j > 2 else (12 if k == 'date' else 18)
    if k in ('commentaire', 'source_detail'): ws.column_dimensions[c.column_letter].width = 28
    if fill and k not in ('date', 'zs', 'commentaire', 'source_detail'):
        for r in range(2, 202): ws.cell(row=r, column=j).fill = PatternFill('solid', fgColor=fill)
ws.row_dimensions[1].height = 78
ws.freeze_panes = 'C2'
for r in range(2, N + 2): ws.cell(row=r, column=1).number_format = 'DD/MM/YYYY'
listes = wb.create_sheet('Listes')
for i, z in enumerate(ZS, 1): listes.cell(row=i, column=1, value=z)
listes.sheet_state = 'hidden'
dv = DataValidation(type='list', formula1=f'=Listes!$A$1:$A$60', allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dv); dv.add(f'B2:B{N + 1}')
dd = DataValidation(type='date', operator='greaterThan', formula1='DATE(2020,1,1)', allow_blank=True, showErrorMessage=True, errorTitle='Date', error='Date au format JJ/MM/AAAA')
ws.add_data_validation(dd); dd.add(f'A2:A{N + 1}')
nb = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True, errorTitle='Nombre', error='Nombre entier positif ou case vide')
ws.add_data_validation(nb); nb.add(f'C2:{ws.cell(row=1, column=len(H) - 2).column_letter}{N + 1}')
wb.save(out); print('écrit', out)
