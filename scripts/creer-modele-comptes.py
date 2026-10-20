"""Génère le classeur à faire remplir pour créer les comptes SurvEpi en une fois.
Usage : python3 scripts/creer-modele-comptes.py [sortie.xlsx] [ZS1,ZS2,...]
Par défaut : ZS du Haut-Uélé, sortie web/src/assets/SurvEpi_modele_comptes.xlsx (téléchargeable depuis l'app)."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

out = sys.argv[1] if len(sys.argv) > 1 else 'web/src/assets/SurvEpi_modele_comptes.xlsx'
ZS = sys.argv[2].split(',') if len(sys.argv) > 2 else ['Isiro', 'Wamba', 'Pawa', 'Boma-Mangbetu', 'Dungu', 'Rungu', 'Gombari', 'Faradje']
ROLES = ['Rapporteur ZS', 'Épidémiologiste superviseur (1re validation ZS)', 'Validateur (2e validation province)',
         'Team Lead / administrateur de la riposte', 'Gestionnaire de données', 'Coordination (lecture seule)', 'MCZ', 'DOS',
         'Administrateur technique (SAV)']
N = 300  # lignes prêtes à remplir

DEEP, BLUE, SOFT = '004E80', '0086C4', 'DFF0F9'
head = Font(bold=True, color='FFFFFF')
headfill = PatternFill('solid', fgColor=DEEP)
reqfill = PatternFill('solid', fgColor='B3302A')
thin = Side(style='thin', color='CFDFEC')
grid = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
lire = wb.active; lire.title = 'À lire'
lignes = [
    ('SurvEpi : création des comptes', 'title'),
    ('', None),
    ('À quoi sert ce fichier', 'h'),
    ('Remplir une ligne par personne qui doit utiliser SurvEpi (onglet « Acteurs »), et compléter les zones de santé (onglet « ZS »).', None),
    ('Le fichier revient ensuite à l’administrateur, qui l’importe dans SurvEpi : tous les comptes sont créés en une fois.', None),
    ('', None),
    ('Onglet « Acteurs »', 'h'),
    ('Colonnes en rouge : obligatoires. Rôle et canal : choisir dans la liste déroulante.', None),
    ('E-mail : obligatoire pour avoir un compte de connexion. Une adresse par personne, jamais partagée.', None),
    ('WhatsApp : 0812345678 ou +243 812 345 678, les deux sont acceptés.', None),
    ('ZS suivies : une ou plusieurs ZS séparées par des virgules (ex. « Isiro, Wamba »). Laisser vide pour les rôles provinciaux.', None),
    ('Rapporteur titulaire de / suppléant de : la ZS dont la personne envoie la fiche du jour. Un seul titulaire et un seul suppléant par ZS.', None),
    ('', None),
    ('Quel rôle choisir ?', 'h'),
    ('Rapporteur ZS : saisit et soumet la fiche journalière de sa ZS.', None),
    ('Épidémiologiste superviseur : 1re validation des fiches de ses ZS.', None),
    ('Validateur : 2e validation, au niveau province (doit être une autre personne que la 1re).', None),
    ('Team Lead / administrateur de la riposte : vue provinciale, validation, réouvertures, paramètres de la riposte.', None),
    ('Gestionnaire de données : contrôle qualité, exports, envoi des WhatsApp préparés.', None),
    ('Coordination : consultation seule.', None),
    ('Administrateur technique (SAV) : comptes, paramètres techniques, sauvegardes.', None),
    ('', None),
    ('Onglet « ZS »', 'h'),
    ('Compléter les cases vides des ZS existantes. Ajouter les nouvelles ZS en bas de la liste.', None),
    ('Dates au format JJ/MM/AAAA. Une case vide ne change rien à ce qui est déjà dans SurvEpi.', None),
    ('', None),
    ('Exemple de ligne « Acteurs »', 'h'),
    ('Jean Mbala | Infirmier titulaire | DPS | Rapporteur ZS | jean.mbala@gmail.com | 0812345678 | Isiro | Isiro | (vide) | WhatsApp', None),
]
for i, (t, kind) in enumerate(lignes, 1):
    c = lire.cell(row=i, column=1, value=t)
    if kind == 'title': c.font = Font(bold=True, size=16, color=DEEP)
    elif kind == 'h': c.font = Font(bold=True, size=12, color=BLUE)
    c.alignment = Alignment(wrap_text=True, vertical='top')
lire.column_dimensions['A'].width = 130
lire.sheet_view.showGridLines = False

listes = wb.create_sheet('Listes')
for i, r in enumerate(ROLES, 1): listes.cell(row=i, column=1, value=r)
for i, c in enumerate(['WhatsApp', 'SMS', 'E-mail', 'Application'], 1): listes.cell(row=i, column=2, value=c)
for i, c in enumerate(['Touchée', 'En préparation'], 1): listes.cell(row=i, column=3, value=c)
for i, c in enumerate(['Oui', 'Non'], 1): listes.cell(row=i, column=4, value=c)
for i, c in enumerate(['Quotidienne', 'Hebdomadaire'], 1): listes.cell(row=i, column=5, value=c)
listes.sheet_state = 'hidden'

def sheet(title, cols):
    ws = wb.create_sheet(title)
    for j, (name, width, req, note) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font = head; c.fill = reqfill if req else headfill
        c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = grid
        if note: c.comment = Comment(note, 'SurvEpi')
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = 'B2'
    return ws

act = sheet('Acteurs', [
    ('Nom complet *', 28, True, None), ('Fonction *', 26, True, 'Ex. Infirmier titulaire, Épidémiologiste OMS'),
    ('Organisation', 16, False, 'Ex. DPS, OMS, MSF'), ('Rôle *', 40, True, 'Choisir dans la liste'),
    ('E-mail', 30, False, 'Obligatoire pour avoir un compte'), ('WhatsApp', 18, False, '0812345678 ou +243812345678'),
    ('ZS suivies', 26, False, 'Plusieurs ZS séparées par des virgules'), ('Rapporteur titulaire de', 22, False, 'Une ZS'),
    ('Rapporteur suppléant de', 22, False, 'Une ZS'), ('Canal préféré', 15, False, 'Pour les rappels'), ('Remarques', 30, False, 'Non importé'),
])
zs = sheet('ZS', [
    ('Zone de santé *', 22, True, None), ('Code DHIS2', 16, False, None), ('Catégorie', 16, False, 'Si la date du 1er cas est inconnue'),
    ('Hotspot', 10, False, None), ('Date du 1er cas confirmé', 16, False, 'JJ/MM/AAAA'), ('Date du dernier cas confirmé', 16, False, 'JJ/MM/AAAA'),
    ('AS total', 10, False, 'Nombre d’aires de santé'), ('AS affectées', 12, False, None), ('MCZ nom', 24, False, None),
    ('Contact MCZ (WhatsApp)', 20, False, None), ('Heure limite', 12, False, 'Vide = heure de la province (16:30)'), ('Fréquence', 15, False, None),
])
for i, n in enumerate(ZS, 2): zs.cell(row=i, column=1, value=n)
for r in range(2, N + 2):
    for c in range(5, 7): zs.cell(row=r, column=c).number_format = 'DD/MM/YYYY'
    zs.cell(row=r, column=11).number_format = 'HH:MM'

def dv(ws, formula, rng, msg):
    v = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True, errorTitle='Valeur non prévue', error=msg)
    ws.add_data_validation(v); v.add(rng)
zsref = f"=ZS!$A$2:$A${N + 1}"
dv(act, f"=Listes!$A$1:$A${len(ROLES)}", f'D2:D{N + 1}', 'Choisir un rôle dans la liste.')
dv(act, '=Listes!$B$1:$B$4', f'J2:J{N + 1}', 'WhatsApp, SMS, E-mail ou Application.')
dv(act, zsref, f'H2:I{N + 1}', 'Choisir une ZS de l’onglet ZS (l’ajouter là si elle manque).')
dv(zs, '=Listes!$C$1:$C$2', f'C2:C{N + 1}', 'Touchée ou En préparation.')
dv(zs, '=Listes!$D$1:$D$2', f'D2:D{N + 1}', 'Oui ou Non.')
dv(zs, '=Listes!$E$1:$E$2', f'L2:L{N + 1}', 'Quotidienne ou Hebdomadaire.')
for ws in (act, zs):
    for row in ws.iter_rows(min_row=2, max_row=N + 1):
        for c in row: c.border = grid
wb.move_sheet('Listes', offset=10)
wb.save(out)
print('écrit', out)
